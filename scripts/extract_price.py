from pathlib import Path

from openpyxl import load_workbook

source = Path('/home/ubuntu/upload/СПКУРАЛ.xlsx')
output = Path('/home/ubuntu/doctor-farmer-quiz/price-extract.md')

workbook = load_workbook(source, data_only=True, read_only=True)
sheet = workbook.active
section = 'НЕ УКАЗАН'
rows = []

for row in sheet.iter_rows(values_only=True):
    values = [str(value).strip() if value is not None else '' for value in row]
    name = values[1] if len(values) > 1 else ''
    cultures = values[7] if len(values) > 7 else ''
    first = values[0] if values else ''
    if first and not name and not cultures:
        section = first
        rows.append(f'## {section}')
        continue
    if name and name != 'Препарат':
        rows.append(f'| {section} | {name} | {values[2]} | {values[3]} | {cultures} |')

output.write_text(
    '# Извлечение номенклатуры из приложенного прайса\n\n'
    '| Раздел | Препарат | Действующее вещество | Норма | Культуры |\n'
    '|---|---|---|---|---|\n' + '\n'.join(rows) + '\n',
    encoding='utf-8'
)
print(output)
